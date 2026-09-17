# -*- coding: utf-8 -*-
"""Mutation-test ladder for verify_rebuttal_r6.py.

For every load-bearing check the R6 verifier makes, inject ONE targeted defect into a COPY of the
letter or package that SHOULD trip that check, run verify_rebuttal_r6.py against the copy, and assert
THE NAMED CHECK goes red. Restore and continue. Disciplines (from the r5 ladder):
  * EVERY MUTATION MUST CHANGE THE INPUT (a no-op raises loudly — a stale test reads as a broken gate).
  * THE BASELINE COPY MUST BE GREEN FIRST.

Run:  PYTHONIOENCODING=utf-8 python scripts/verify_rebuttal_r6_mutations.py
"""
# --- portable roots -----------------------------------------------------------------------------
# Injected by scripts/build_repo.py. The working tree hardcoded an absolute local path; the deposit
# resolves it from CKM_P4_BASE, or from this file's own location (repo/code/ plays the role the
# working tree called independent_build/). Raw GWAS summary statistics are NOT deposited: set
# CKM_ROOT to wherever you obtained them if you intend to re-run the upstream extraction steps.
import os as _os
P4_BASE = _os.environ.get("CKM_P4_BASE") or _os.path.dirname(
    _os.path.dirname(_os.path.abspath(__file__)))
P4_ROOT = _os.environ.get("CKM_P4_ROOT") or _os.path.dirname(P4_BASE)
CKM_ROOT = _os.environ.get("CKM_ROOT") or _os.path.dirname(P4_ROOT)
SHARED_LIB = _os.environ.get("CKM_SHARED_LIB") or _os.path.join(CKM_ROOT, "_shared")
# ------------------------------------------------------------------------------------------------

import io, json, os, re, sys, shutil, tempfile, subprocess

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
sys.path.insert(0, SHARED_LIB)
from pdftext import pdf_text                                                # noqa: E402

ROOT = P4_ROOT
BASE = os.path.join(ROOT, "independent_build")
SCR = os.path.join(BASE, "scripts")
_bsv = io.open(os.path.join(SCR, "build_submission_v1.py"), encoding="utf-8").read()
PKG_VERSION = re.search(r'^PKG_VERSION\s*=\s*"([^"]+)"', _bsv, flags=re.M).group(1)
R_PKG = os.path.join(ROOT, f"submission package {PKG_VERSION}")
R_LETTER = os.path.join(BASE, "RESPONSE_TO_INTERNAL_REVIEW_R6.md")

WORK = tempfile.mkdtemp(prefix="r6mut_")
TMP_PKG = os.path.join(WORK, "pkg")
TMP_LETTER = os.path.join(WORK, "letter.md")
shutil.copytree(R_PKG, TMP_PKG)
shutil.copy(R_LETTER, TMP_LETTER)


def restore(rel):
    if rel == "letter":
        shutil.copy(R_LETTER, TMP_LETTER)
    else:
        shutil.copy(os.path.join(R_PKG, rel), os.path.join(TMP_PKG, rel))


def _sub(path, find, repl, count=0):
    t = io.open(path, encoding="utf-8").read()
    t2 = re.sub(find, repl, t, count=count)
    assert t2 != t, f"NO-OP mutation (target gone?): {find!r} in {os.path.basename(path)}"
    io.open(path, "w", encoding="utf-8").write(t2)


def L(find, repl, count=0):
    _sub(TMP_LETTER, find, repl, count)


def M(rel, find, repl, count=0):
    _sub(os.path.join(TMP_PKG, rel), find, repl, count)


def M_append(rel, text):
    io.open(os.path.join(TMP_PKG, rel), "a", encoding="utf-8").write(text)


def PDF_add(rel, text):
    import fitz
    p = os.path.join(TMP_PKG, rel)
    d = fitz.open(p)
    d[0].insert_text(fitz.Point(72, 72), text, fontsize=11)
    d.save(p, incremental=True, encryption=fitz.PDF_ENCRYPT_KEEP)
    d.close()
    assert text.split()[0] in pdf_text(p), f"PDF add did not take in {rel}"


def PDF_swap(rel, src_rel):
    shutil.copy(os.path.join(TMP_PKG, src_rel), os.path.join(TMP_PKG, rel))


def _strip_interim():
    # The interim wording is intentionally present in several declaration surfaces; to trip the
    # scope-wide presence check it must be removed from every one (a single-file edit is a no-op MISS).
    import docx
    hit = False
    for rel in ("02_cover_declarations/DECLARATIONS.md", "02_cover_declarations/COVER_LETTER.md"):
        p = os.path.join(TMP_PKG, rel)
        t = io.open(p, encoding="utf-8").read()
        if "provided prior to publication" in t:
            io.open(p, "w", encoding="utf-8").write(t.replace("provided prior to publication", "removed"))
            hit = True
    for rel in ("02_cover_declarations/DECLARATIONS.docx", "02_cover_declarations/COVER_LETTER.docx"):
        p = os.path.join(TMP_PKG, rel)
        d = docx.Document(p)
        for para in d.paragraphs:
            if "provided prior to publication" in para.text:
                new = para.text.replace("provided prior to publication", "removed")
                for r in para.runs:
                    r.text = ""
                (para.runs[0] if para.runs else para.add_run("")).text = new
                hit = True
        d.save(p)
    assert hit, "NO-OP: interim wording not found to strip"


# ---------- the ladder: (expected-failing check label, [files to restore], apply) ----------------
# ---- targets DERIVED from the live package, never typed ------------------------------------------
# Literal targets ("(99 files)", "main text **10,617**", "**52 references**") went stale as the package
# moved and made the ladder refuse. These are read at run time so they cannot drift.
_CF = json.load(io.open(os.path.join(BASE, "manifest", "CANONICAL_FACTS.json"), encoding="utf-8"))
_MAIN, _REFS = _CF["main_text_words"], _CF["references_cited"]
_NFILES = sum(len(_f) for _r, _d, _f in os.walk(R_PKG))
_DEP = json.load(io.open(os.path.join(BASE, "manifest", "DEPOSIT.json"), encoding="utf-8"))

MUT = [
    # ---- §1 letter counts vs derive_facts ----
    ("§1 main-text words match derive_facts", ["letter"],
     lambda: L(rf"main text \*\*{_MAIN:,}\*\*", f"main text **{_MAIN - 17:,}**")),
    ("§1 reference count matches", ["letter"],
     lambda: L(rf"\*\*{_REFS} references\*\*", f"**{_REFS - 1} references**")),
    # ---- §2 M1 graphical abstract ----
    ("§2 M1: the graphical-abstract PDF prints +0.42", ["03_main_figures/GraphicalAbstract.pdf"],
     lambda: PDF_swap("03_main_figures/GraphicalAbstract.pdf", "03_main_figures/Figure1.pdf")),
    ("§2 M1: the stale +0.41 is gone from the graphical-abstract PDF",
     ["03_main_figures/GraphicalAbstract.pdf"],
     lambda: PDF_add("03_main_figures/GraphicalAbstract.pdf", "CAD-INDEPENDENT +0.41 log-odds")),
    # ---- §2 M2 placeholders ----
    ("§2 M2: no bracketed author-supplied / repository placeholder survives anywhere",
     ["02_cover_declarations/DECLARATIONS.md"],
     lambda: M_append("02_cover_declarations/DECLARATIONS.md", "\n[GitHub URL — AUTHOR-SUPPLIED]\n")),
    ("§2 M2: the title-page typesetting instruction is removed",
     ["02_cover_declarations/TITLE_PAGE.md"],
     lambda: M_append("02_cover_declarations/TITLE_PAGE.md",
                      "\nAffiliation numbers are to be set as superscripts at typesetting.\n")),
    # the defect that actually happened: the VERSION DOI pasted where the CONCEPT DOI belongs
    ("§2 M2: the repository identifiers are now the real, archived ones (concept DOI + URL)",
     ["02_cover_declarations/DECLARATIONS.md"],
     lambda: M("02_cover_declarations/DECLARATIONS.md", _DEP["concept_doi"], _DEP["version_doi"])),
    # ---- §2 M3 LLM disclosure ----
    ("§2 M3: the LLM disclosure names Anthropic Claude", ["01_manuscript/METHODS.md"],
     lambda: M("01_manuscript/METHODS.md", "Anthropic Claude", "an assistant")),
    ("§2 M3: the author-facing confirmation instruction is removed", ["01_manuscript/METHODS.md"],
     lambda: M_append("01_manuscript/METHODS.md", "\nconfirm that this statement covers every tool used.\n")),
    ("§2 M3: the disclosure closes with the raw-data / autonomy statement", ["01_manuscript/METHODS.md"],
     lambda: M("01_manuscript/METHODS.md", "No AI system accessed raw data", "The system did things")),
    # ---- §2 M4 revision-history language ----
    ("§2 M4: no revision-history phrase remains in the manuscript sections", ["01_manuscript/DISCUSSION.md"],
     lambda: M_append("01_manuscript/DISCUSSION.md", "\nThis survived the crossing into East Asians.\n")),
    # ---- §3 M5 references ----
    ("§3 M5: Reference 4 (Zheng) keeps the 2022 issue year", ["01_manuscript/REFERENCES_NUMBERED.md"],
     lambda: M("01_manuscript/REFERENCES_NUMBERED.md",
               r"Int J Epidemiol\. 2022;50\(6\):1995-2010", "Int J Epidemiol. 2021;50(6):1995-2010")),
    ("§3 M5: Reference 5 (Yin) carries the article number 189", ["01_manuscript/REFERENCES_NUMBERED.md"],
     lambda: M("01_manuscript/REFERENCES_NUMBERED.md",
               r"Cardiovasc Diabetol\. 2026;25:189", "Cardiovasc Diabetol. 2026;25(1)")),
    ("§3 M5: the TPMI cohort reference (Yang, PMID 41092961) is cited",
     ["01_manuscript/REFERENCES_NUMBERED.md"],
     lambda: M("01_manuscript/REFERENCES_NUMBERED.md", "41092961", "49999999")),
    ("§3 M5: the Taiwan Biobank cohort reference (Wei, PMID 33574314) is cited",
     ["01_manuscript/REFERENCES_NUMBERED.md"],
     lambda: M("01_manuscript/REFERENCES_NUMBERED.md", "33574314", "39999999")),
    # ---- §4 figures + dialect ----
    ("§4 Figure 2 panel d is relabelled 'Q-minimisation' (shorthand 'qhet' gone)",
     ["03_main_figures/Figure2.pdf"],
     lambda: PDF_swap("03_main_figures/Figure2.pdf", "03_main_figures/Figure1.pdf")),
    ("§4 Figure 3 uses the contiguous 'correlated-marker negative control' label",
     ["03_main_figures/Figure3.pdf"],
     lambda: PDF_add("03_main_figures/Figure3.pdf", "an unqualified negative control here")),
    ("§4 the metabolic dialect / hyphenation is harmonised to British in the manuscript",
     ["01_manuscript/RESULTS.md"],
     lambda: M_append("01_manuscript/RESULTS.md", "\nWe observed dysglycemia in the cohort.\n")),
]


def run_verifier():
    env = dict(os.environ, CKM_R6_LETTER=TMP_LETTER, CKM_R6_PKG=TMP_PKG, PYTHONIOENCODING="utf-8")
    r = subprocess.run([sys.executable, os.path.join(SCR, "verify_rebuttal_r6.py")],
                       capture_output=True, text=True, encoding="utf-8", errors="replace", env=env)
    return (r.stdout or "") + (r.stderr or "")


def failed_labels(out):
    return set(re.findall(r"✗ FAILED: (.+?)(?:  \[|$)", out, flags=re.M))


TRIPPED, MISSED = [], []
base = run_verifier()
_bl = re.search(r"RESULT: (\d+) verified, (\d+) FAILED", base)
assert _bl and _bl.group(2) == "0", "BASELINE COPY IS NOT GREEN:\n" + base[-1500:]
print(f"baseline: {_bl.group(1)} verified, 0 FAILED (clean copy)\n")

for label, files, apply in MUT:
    for f in files:
        restore(f)
    apply()
    fl = failed_labels(run_verifier())
    if label in fl:
        TRIPPED.append(label)
        print(f"  ✓ tripped: {label}")
    else:
        MISSED.append(label)
        print(f"  ✗ MISSED (check is vacuous or mutation stale): {label}\n       failed instead: {sorted(fl)[:3]}")
    for f in files:
        restore(f)

print(f"\nLADDER: {len(TRIPPED)}/{len(MUT)} checks tripped by their named defect")
if MISSED:
    print("MISSED:")
    for m in MISSED:
        print("  -", m)
shutil.rmtree(WORK, ignore_errors=True)
sys.exit(1 if MISSED else 0)
