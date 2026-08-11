# -*- coding: utf-8 -*-
"""Derive every duplicated fact about CKM Paper 4 from its source files, once.

Nothing here is hand-typed: word counts, panel counts, figure/table/reference counts are all
computed from the canonical sources. The builder and the gate script both consume the JSON this
writes, so a number can never drift between the title page, the README and the manuscript.

Word-count convention (reverse-engineered to reproduce the recorded section counts exactly):
strip HTML comments, drop markdown heading lines, count whitespace-separated tokens.
Abstract convention additionally drops standalone math-operator tokens.
"""
# --- portable roots -----------------------------------------------------------------------------
# Injected by scripts/build_repo.py. The working tree hardcoded an absolute local path; the deposit
# resolves it from CKM_P4_BASE, or from this file's own location (repo/code/ plays the role the
# working tree called independent_build/). Raw GWAS summary statistics are NOT deposited: set
# CKM_ROOT to wherever you obtained them if you intend to re-run the upstream extraction steps.
import os as _os
P4_BASE = _os.environ.get("CKM_P4_BASE") or _os.path.dirname(
    _os.path.dirname(_os.path.abspath(__file__)))
P4_ROOT = _os.environ.get("CKM_P4_ROOT") or _os.path.dirname(P4_BASE)
CKM_ROOT = _os.environ.get("CKM_ROOT") or _os.path.dirname(P4_ROOT)
SHARED_LIB = _os.environ.get("CKM_SHARED_LIB") or _os.path.join(CKM_ROOT, "_shared")
# ------------------------------------------------------------------------------------------------

import os, re, json, subprocess, sys

BASE = P4_BASE
FIG = os.path.join(BASE, "figures")
MAN = os.path.join(BASE, "manuscript")
RES = os.path.join(BASE, "results")

def rd(p):
    return open(p, encoding="utf-8").read() if os.path.exists(p) else ""

def strip_comments(t):
    return re.sub(r"<!--.*?-->", "", t, flags=re.S)

def wc_section(t):
    """Section word count: no HTML comments, no heading lines."""
    t = strip_comments(t)
    t = re.sub(r"^#.*$", "", t, flags=re.M)
    return len(t.split())

OPERATORS = {"=", "×", "+", "−", "-", "<", ">", "≤", "≥", "~", "±", "/"}

def wc_abstract(t):
    """Abstract word count: journals do not count standalone math operators."""
    t = strip_comments(t)
    t = re.sub(r"^#.*$", "", t, flags=re.M)
    t = re.sub(r"^_.*_$", "", t, flags=re.M)          # italic meta lines
    t = t.replace("**", "").replace("---", " ")
    return len([w for w in t.split() if w not in OPERATORS])

facts = {}

# ---------- section word counts ----------
sections = {}
for name in ["INTRODUCTION", "METHODS", "RESULTS", "DISCUSSION", "CONCLUSIONS"]:
    sections[name] = wc_section(rd(os.path.join(FIG, f"{name}.md")))
facts["word_counts"] = sections
facts["main_text_words"] = sum(sections.values())

# ---------- abstract variants ----------
ab = strip_comments(rd(os.path.join(FIG, "ABSTRACT.md")))
def slice_between(t, start_pat, end_pat):
    m = re.search(start_pat, t)
    if not m:
        return ""
    rest = t[m.end():]
    e = re.search(end_pat, rest)
    return rest[:e.start()] if e else rest

full = slice_between(ab, r"##\s*FULL abstract[^\n]*\n", r"\n---")
cap = slice_between(ab, r"##\s*CAP variant[^\n]*\n", r"\n---")
facts["abstract_full_words"] = wc_abstract(full)
facts["abstract_cap_words"] = wc_abstract(cap)

# ---------- main figures + panels ----------
leg = rd(os.path.join(FIG, "FIGURE_LEGENDS.md"))
panels = {}
for m in re.finditer(r"\*\*Figure (\d+) \|", leg):
    n = m.group(1)
    nxt = leg.find("**Figure", m.end())
    seg = leg[m.end(): nxt if nxt > 0 else len(leg)]
    panels[n] = sorted(set(re.findall(r"\(\*{0,2}([a-z])\*{0,2}\)", seg)))
facts["main_figures"] = len(panels)
facts["panels_per_figure"] = {k: len(v) for k, v in sorted(panels.items())}
facts["total_panels"] = sum(len(v) for v in panels.values())

# figure files actually on disk
facts["main_figure_files"] = sorted(
    os.path.basename(p) for p in
    [os.path.join(FIG, f"Figure{i}.{e}") for i in range(1, 20) for e in ("png", "pdf")]
    if os.path.exists(p))

# ---------- supplementary figures ----------
sf = rd(os.path.join(FIG, "SUPPL_FIGURES.md"))
# Anchor at line start: a legend HEADER only. A bare "**Supplementary Figure S6**" also appears mid-line
# in the italic note recording S6's retirement, and must not be counted as a defined figure.
facts["supp_figures_defined"] = sorted(set(
    int(x) for x in re.findall(r"^\*\*Supplementary Figure S(\d+) \|", sf, flags=re.M)))
facts["supp_figure_files"] = sorted(
    os.path.basename(p) for p in
    [os.path.join(FIG, f"SupplFig{i}.{e}") for i in range(1, 20) for e in ("png", "pdf")]
    if os.path.exists(p))
facts["supp_figures"] = len(facts["supp_figures_defined"])

# ---------- supplementary tables ----------
try:
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(MAN, "Supplementary_Tables.xlsx"), read_only=True)
    facts["supp_table_sheets"] = len(wb.sheetnames)
    facts["supp_table_sheet_names"] = list(wb.sheetnames)
    ids = sorted(set(int(x) for s in wb.sheetnames for x in re.findall(r"S(\d+)", s)))
    facts["supp_tables"] = len(ids)
    facts["supp_table_ids"] = ids
    wb.close()
except Exception as e:
    facts["supp_table_error"] = str(e)

# ---------- references ----------
r = subprocess.run([sys.executable, os.path.join(BASE, "scripts", "build_numbered_refs.py")],
                   capture_output=True, text=True)
m = re.search(r"(\d+) cited, (\d+) uncited, (\d+) unmatched", r.stdout)
if m:
    facts["references_cited"] = int(m.group(1))
    facts["references_uncited"] = int(m.group(2))
    facts["references_unmatched"] = int(m.group(3))
else:
    facts["references_error"] = (r.stdout + r.stderr)[-500:]

# ---------- network scale (from results, not prose) ----------
import csv
# The network SIZE is the number of directed edges tested = rows of forward_local_edges.csv (132
# since the network was completed; this is READ from the file, never typed, so it tracks automatically).
# edges_standardised.csv is a smaller, scale-standardised subset (101) and is NOT the network size;
# conflating the two would understate the Bonferroni denominator quoted throughout the Methods.
for key, fname in [("network_edges_total", "forward_local_edges.csv"),
                   ("network_edges_standardised", "edges_standardised.csv")]:
    p = os.path.join(RES, fname)
    if os.path.exists(p):
        with open(p, encoding="utf-8") as fh:
            facts[key] = len(list(csv.DictReader(fh)))

# Bonferroni threshold quoted in the Methods, recomputed rather than trusted.
if facts.get("network_edges_total"):
    facts["bonferroni_threshold"] = 0.05 / facts["network_edges_total"]
    with open(os.path.join(RES, "forward_local_edges.csv"), encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    facts["edges_passing_bonferroni"] = sum(
        1 for r in rows if r.get("ivw_p") and float(r["ivw_p"]) < facts["bonferroni_threshold"])

if __name__ == "__main__":
    out = os.path.join(BASE, "manifest", "CANONICAL_FACTS.json")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    json.dump(facts, open(out, "w", encoding="utf-8"), indent=2, ensure_ascii=False)
    for k, v in facts.items():
        print(f"{k:26s} = {v}")
    print("\nwrote", os.path.relpath(out, BASE))
