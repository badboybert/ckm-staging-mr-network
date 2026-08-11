# -*- coding: utf-8 -*-
"""Check every FACTUAL claim the round-3 review makes about the shipped package.

The review is an input, not a specification. Some of its findings may already be fixed (v3.2 was
rebuilt several times on the same day the review was written), some may be true, and some may be
about surfaces it could not open. Triage must be built on what the artifact actually says, so each
checkable claim is restated here as a search over every shipped surface: .md, both .docx twins, the
workbook cell text, and the text layer of all 12 figure PDFs.

Prints CONFIRMED (the review is right, this is real work) or STALE (already fixed) per claim.
"""
# --- portable roots -----------------------------------------------------------------------------
# Injected by scripts/build_repo.py. The working tree hardcoded an absolute local path; the deposit
# resolves it from CKM_P4_BASE, or from this file's own location (repo/code/ plays the role the
# working tree called independent_build/). Raw GWAS summary statistics are NOT deposited: set
# CKM_ROOT to wherever you obtained them if you intend to re-run the upstream extraction steps.
import os as _os
P4_BASE = _os.environ.get("CKM_P4_BASE") or _os.path.dirname(
    _os.path.dirname(_os.path.abspath(__file__)))
P4_ROOT = _os.environ.get("CKM_P4_ROOT") or _os.path.dirname(P4_BASE)
CKM_ROOT = _os.environ.get("CKM_ROOT") or _os.path.dirname(P4_ROOT)
SHARED_LIB = _os.environ.get("CKM_SHARED_LIB") or _os.path.join(CKM_ROOT, "_shared")
# ------------------------------------------------------------------------------------------------

import io, os, re, sys, glob, zlib

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

ROOT = P4_ROOT
# 2026-07-26: the package version is DERIVED from build_submission_v1.py, never typed here. A
# checker hard-pointed at an old package silently verifies an artifact nobody is shipping - the
# exact failure prose_ai_scan.py had. Its scope IS the check.
_bsv = io.open(os.path.join(ROOT, "independent_build", "scripts", "build_submission_v1.py"),
               encoding="utf-8").read()
_m = re.search(r'^PKG_VERSION\s*=\s*"([^"]+)"', _bsv, flags=re.M)
assert _m, "could not read PKG_VERSION from build_submission_v1.py"
PKG_VERSION = _m.group(1)
PKG = os.path.join(ROOT, f"submission package {PKG_VERSION}")
BASE = os.path.join(ROOT, "independent_build")


def docx_text(p):
    import docx
    d = docx.Document(p)
    return "\n".join([q.text for q in d.paragraphs] +
                     [c.text for t in d.tables for r in t.rows for c in r.cells])


# PDF text comes from the SHARED reader (~/.claude/lib/pdftext.py), not a local copy. This file
# carried its own scanner until 2026-07-27; `pdftext.py --audit` found five such forks in this
# project, and measuring them showed the local scanner returned ~70x too many characters — CMap and
# font-encoding junk — so every presence probe over it was weaker than it looked. The import raises
# if the shared library is missing: a checker that cannot read must never report clean.
sys.path.insert(0, SHARED_LIB)
from pdftext import pdf_text
SURF = {}
for p in glob.glob(os.path.join(PKG, "**", "*.md"), recursive=True):
    SURF[os.path.relpath(p, PKG)] = open(p, encoding="utf-8").read()
for p in glob.glob(os.path.join(PKG, "**", "*.docx"), recursive=True):
    SURF[os.path.relpath(p, PKG)] = docx_text(p)
for p in glob.glob(os.path.join(PKG, "**", "*.pdf"), recursive=True):
    SURF["PDF:" + os.path.basename(p)] = pdf_text(p)
import openpyxl
_wb = openpyxl.load_workbook(os.path.join(PKG, "04_supplementary", "Supplementary_Tables.xlsx"),
                             read_only=True, data_only=True)
SURF["WORKBOOK"] = "\n".join(str(v) for ws in _wb.worksheets
                             for row in ws.iter_rows(values_only=True)
                             for v in row if isinstance(v, str))

CONFIRMED, STALE = [], []


def claim(tag, pattern, review_says_present=True, scope=None, note=""):
    """review_says_present: the review asserts this string IS in the package."""
    pool = SURF if scope is None else {k: v for k, v in SURF.items() if scope in k}
    hits = sorted({k for k, t in pool.items() if re.search(pattern, t, re.I)})
    real = bool(hits) if review_says_present else (not hits)
    rec = f"{tag}\n      where: {hits[:4] if hits else 'nowhere'}" + (f"\n      note: {note}" if note else "")
    (CONFIRMED if real else STALE).append(rec)


print("=" * 96)
print(f"ROUND-3 CLAIMS CHECKED AGAINST submission package {PKG_VERSION}")
print(f"surfaces opened: {len(SURF)}")
print("=" * 96)

# ---- Table 3, the rebuttal-to-package consistency audit -----------------------------------------
claim("A1  workbook/SupplFig S7 still state the 68-79% mediated fraction", r"68\s*[-–]\s*79")
claim("A2  SupplFig S7 still labels staging HIGHER CONF.",
      r"AHA staging[^|]{0,40}HIGHER CONF|HIGHER CONF[^|]{0,40}staging", scope="SupplFig7")
claim("A3  SupplFig S7 says 'not a selection artifact'", r"not a selection artifact")
claim("A4a Figure 3e uses 'true causal' / 'true positive' and 'false positive'",
      r"true positive|true causal|false positive", scope="PDF:Figure3")
# "causal" now legitimately appears as "causal model favoured" and "causal-model preference
# threshold" -- that IS the fix. Match only the bare binary verdict the review objected to.
claim("A4b Figure 3b renders CAUSE as a BINARY causal verdict",
      r"\bcausal\b(?!\s*(model|-model))(?![a-z-])", scope="PDF:Figure3")
claim("A5a Figure 4c foregrounds an identity-line raw-magnitude comparison",
      r"identity", scope="PDF:Figure4")
claim("A5b Results still foregrounds effect magnitudes 'tracking the identity line'",
      r"track the identity line|effect sizes track")
claim("A6  Methods retain position-only primary SBP extraction",
      r"matched to the SBP file by position alone")
claim("A7  Manuscript still says MR-PRESSO was targeted rather than all 59",
      r"MR-PRESSO and CAUSE were applied to the load-bearing")
claim("A8  no complete Steiger-margin table in the shipped workbook",
      r"margin", scope="WORKBOOK", review_says_present=False,
      note="review says the margin table is NOT evident; a hit here means it IS present")

# ---- diction the review asks to remove ----------------------------------------------------------
print("\n-- Section 6.3 diction list, checked over every shipped surface --")
for word, pat in [
    ("'direct driver'", r"direct driver"),
    ("'predominantly via/through CAD' for T2D", r"predominantly,? (not exclusively,? )?through coronary|predominantly via CAD"),
    ("'cleanly' / 'inserts cleanly'", r"inserts cleanly|resolve[sd]? cleanly|cleanly"),
    ("'genuine null'", r"genuine null"),
    ("'true causal/false positive'", r"true causal|false positive"),
    ("'not a selection artifact'", r"not a selection artifact"),
    ("'pan-subtype'", r"pan-subtype"),
    ("'track the identity line'", r"track the identity line"),
    ("'vicious cycles'", r"vicious cycle"),
    ("'load-bearing'", r"load-bearing"),
    ("'earned that reading'", r"earned that reading"),
    ("'hardened the hedge'", r"hardened the hedge"),
    ("'knife-edge'", r"knife-edge"),
    ("'fully mediated' for >100%", r"fully mediated"),
    ("'not sig (power)'", r"not sig \(power\)|not significant \(power\)"),
    # bare "overturned" matched "discordance would have overturned the order", which is ordinary
    # English about falsifiability, not the retired Figure 4e label.
    ("'BBJ null overturned'", r"null overturned|overturned by TPMI"),
    ("'Node-level bootstrap'", r"[Nn]ode-level bootstrap"),
]:
    hits = sorted({k for k, t in SURF.items() if re.search(pat, t, re.I)})
    print(f"  {'PRESENT' if hits else 'ABSENT ':8s} {word:42s} {hits[:3]}")

# ---- the degree-preserving null implementation --------------------------------------------------
print("\n-- BLOCKER: degree-preserving null mixing (read from the code, not the review) --")
src = ""
for f in glob.glob(os.path.join(BASE, "scripts", "*.py")):
    t = open(f, encoding="utf-8", errors="replace").read()
    if "degree" in t.lower() and ("rewir" in t.lower() or "switch" in t.lower()):
        src = f
        for m in re.finditer(r"^.*(n_switch|switch|attempt|rewir|nswap|swap).*$", t, flags=re.M | re.I):
            line = m.group(0).strip()
            if re.search(r"=|range\(|for |def ", line) and len(line) < 160:
                print(f"    {os.path.basename(f)}: {line}")
        break
if not src:
    print("    (no degree-preserving rewiring implementation found by keyword)")

# ---- Research Insights format -------------------------------------------------------------------
print("\n-- MAJOR: Research Insights format --")
ri = SURF.get("01_manuscript\\RESEARCH_INSIGHTS.md") or SURF.get("01_manuscript/RESEARCH_INSIGHTS.md", "")
# Measure the DELIVERABLE (the .docx table the journal receives), not the markdown source cell.
# The source separates highlights with <br>; counting the raw cell reported a 265-character
# "paragraph" for a table that renders as three compliant lines.
import docx as _docx
_d = _docx.Document(os.path.join(PKG, "01_manuscript", "CKM_Paper4_manuscript.docx"))
for _r in _d.tables[0].rows:
    _q = _r.cells[0].text.strip()
    _lines = [x.text.strip() for x in _r.cells[1].paragraphs if x.text.strip()]
    _lens = [len(x) for x in _lines]
    _ok = all(n <= 100 for n in _lens)
    print(f"    [{len(_lines)} separate highlight(s), longest {max(_lens)} chars] "
          f"{'OK' if _ok else 'OVER 100 CHARS'}  {_q[:44]}")

print("\n" + "=" * 96)
print(f"CONFIRMED (real, still present in the shipped package): {len(CONFIRMED)}")
for c in CONFIRMED:
    print("  [X] " + c)
print(f"\nSTALE (review is describing an earlier build): {len(STALE)}")
for c in STALE:
    print("  [ok] " + c)
print("=" * 96)
