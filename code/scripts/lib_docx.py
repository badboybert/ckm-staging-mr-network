# -*- coding: utf-8 -*-
"""One markdown-table renderer for every .docx this project builds.

WHY THIS FILE EXISTS
--------------------
Round 4 found the Additional-files table shipping in CKM_Paper4_manuscript.docx as SEVEN paragraphs
of literal pipe syntax, separator row included:

    | File | Format | Title | Description |
    |---|---|---|---|
    | Additional file 1 | DOCX | STROBE-MR checklist | ... |

The same table renders correctly in DECLARATIONS.docx, because build_submission_v1.py has a
`render_table` and build_manuscript_docx.py's `add_body` never had one. Two renderers, one of them
fixed — the fork pattern in CLAUDE.md §9. Both now import from here, so a fix cannot land in one and
not the other.

`add_inline` is passed in rather than imported: the two builders style runs differently (font size,
bold headers) and only the TABLE construction is shared.
"""
# --- portable roots -----------------------------------------------------------------------------
# Injected by scripts/build_repo.py. The working tree hardcoded an absolute local path; the deposit
# resolves it from CKM_P4_BASE, or from this file's own location (repo/code/ plays the role the
# working tree called independent_build/). Raw GWAS summary statistics are NOT deposited: set
# CKM_ROOT to wherever you obtained them if you intend to re-run the upstream extraction steps.
import os as _os
P4_BASE = _os.environ.get("CKM_P4_BASE") or _os.path.dirname(
    _os.path.dirname(_os.path.abspath(__file__)))
P4_ROOT = _os.environ.get("CKM_P4_ROOT") or _os.path.dirname(P4_BASE)
CKM_ROOT = _os.environ.get("CKM_ROOT") or _os.path.dirname(P4_ROOT)
SHARED_LIB = _os.environ.get("CKM_SHARED_LIB") or _os.path.join(CKM_ROOT, "_shared")
# ------------------------------------------------------------------------------------------------

from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SEP = set("-: ")


def is_table_row(line):
    """A markdown table row: starts and ends with a pipe and has at least one interior pipe."""
    s = line.strip()
    return s.startswith("|") and s.endswith("|") and s.count("|") >= 3


def split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def _no_split(row):
    """Word: <w:cantSplit/> — do not break this row across a page boundary."""
    pr = row._tr.get_or_add_trPr()
    if pr.find(qn("w:cantSplit")) is None:
        pr.append(OxmlElement("w:cantSplit"))


def _repeat_header(row):
    """Word: <w:tblHeader/> — repeat this row at the top of every page the table spans."""
    pr = row._tr.get_or_add_trPr()
    if pr.find(qn("w:tblHeader")) is None:
        pr.append(OxmlElement("w:tblHeader"))


def render_md_table(doc, rows, add_inline, font_pt=8, bold_header=True):
    """Render consecutive markdown pipe rows as a real Word table. Drops the |---|---| separator.

    Round 5 (M6): every row is marked cantSplit and the header row tblHeader. The Additional-files
    table was breaking mid-row across manuscript pages 44-45, leaving the orphan fragments
    "SupplFig7)." and "study." stranded at the top of the next page — in the manuscript AND in the
    standalone Declarations, because both render through this one function. Setting it here means
    neither builder can regress, and any future table inherits the same behaviour.
    """
    cells = [split_row(r) for r in rows]
    cells = [c for c in cells if not all(set(x) <= SEP for x in c)]
    if not cells:
        return None
    ncol = len(cells[0])
    t = doc.add_table(rows=len(cells), cols=ncol)
    t.style = "Table Grid"
    for ri, row in enumerate(t.rows):
        _no_split(row)
        if ri == 0 and bold_header:
            _repeat_header(row)
    for ri, row in enumerate(cells):
        for ci, val in enumerate(row):
            if ci >= ncol:
                continue
            cell = t.cell(ri, ci)
            cell.paragraphs[0].text = ""
            add_inline(cell.paragraphs[0], val)
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(font_pt)
                if ri == 0 and bold_header:
                    run.font.bold = True
    return t
