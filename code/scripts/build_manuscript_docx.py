# -*- coding: utf-8 -*-
"""Assemble the CKM Paper 4 manuscript into a single .docx from the verified section files.

The section sources are authored in AUTHOR-DATE because that is the form a human can edit without
renumbering anything. The target journal (Cardiovascular Diabetology) requires numbered Vancouver,
so the conversion is applied HERE, at build time, by scripts/lib_vancouver.py -- the same transform
build_submission_v1.py and gate_package.py apply, so the .docx, the shipped .md and the gate's
expectation are the same function of the same source and cannot drift apart.
Author/creator metadata = Bertrand Chin-Ming Tan (global rule 8)."""
# --- portable roots -----------------------------------------------------------------------------
# Injected by scripts/build_repo.py. The working tree hardcoded an absolute local path; the deposit
# resolves it from CKM_P4_BASE, or from this file's own location (repo/code/ plays the role the
# working tree called independent_build/). Raw GWAS summary statistics are NOT deposited: set
# CKM_ROOT to wherever you obtained them if you intend to re-run the upstream extraction steps.
import os as _os
P4_BASE = _os.environ.get("CKM_P4_BASE") or _os.path.dirname(
    _os.path.dirname(_os.path.abspath(__file__)))
P4_ROOT = _os.environ.get("CKM_P4_ROOT") or _os.path.dirname(P4_BASE)
CKM_ROOT = _os.environ.get("CKM_ROOT") or _os.path.dirname(P4_ROOT)
SHARED_LIB = _os.environ.get("CKM_SHARED_LIB") or _os.path.join(CKM_ROOT, "_shared")
# ------------------------------------------------------------------------------------------------

import os, re, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import lib_vancouver as VAN
import lib_docx
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

FIG  = os.path.join(P4_BASE, "figures")
OUT  = os.path.join(P4_BASE, "manuscript")
os.makedirs(OUT, exist_ok=True)

def strip_comment(t):
    return re.sub(r"<!--.*?-->", "", t, flags=re.S).strip()

MASTER, REF_ORDER, AMAP = VAN.numbering()

def read(name):
    """Read a section and, for the sections that carry citations, convert author-date to Vancouver."""
    t = strip_comment(open(os.path.join(FIG, name), encoding="utf-8").read())
    if name in VAN.REWRITE:
        src = t
        t = VAN.to_vancouver(t, AMAP, MASTER)
        left = VAN.unconverted(t, MASTER, AMAP)
        assert not left, f"{name}: citations left unconverted -> {left[:5]}"
        bd = VAN.bracket_delta(src, t)
        assert bd["("] == bd[")"] and bd["["] == bd["]"], f"{name}: rewrite unbalanced brackets {bd}"
    return t

# --- inline **bold** / *italic* -> runs ---
# ^...^ is a superscript run (the affiliation markers on the title page). The packager's renderer
# has understood this since round 6; this one did not, so deriving the author line from
# FRONT_MATTER.md (round-7 C002) would otherwise print literal carets in the manuscript.
INLINE = re.compile(r"\*\*(.+?)\*\*|\*(.+?)\*|\^([^\^\n]+)\^")
def add_inline(p, text):
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            p.add_run(text[pos:m.start()])
        if m.group(1) is not None:
            p.add_run(m.group(1)).bold = True
        elif m.group(2) is not None:
            p.add_run(m.group(2)).italic = True
        else:
            p.add_run(m.group(3)).font.superscript = True
        pos = m.end()
    if pos < len(text):
        p.add_run(text[pos:])

def add_body(doc, md, h1_title=None):
    """Render a section .md: '# X' -> Heading 1 (overridable), '## X' -> Heading 2, else paragraph.

    Markdown pipe tables become real Word tables. Without this branch they shipped as literal body
    text — the Additional-files block reached the submitted manuscript as seven paragraphs including
    the raw `|---|---|---|---|` separator, while the SAME table rendered correctly in
    DECLARATIONS.docx because the packager had a renderer and this did not (scripts/lib_docx.py).
    """
    lines = md.split("\n")
    i, n = 0, len(lines)
    while i < n:
        s = lines[i].rstrip()
        if not s:
            i += 1
            continue
        if re.match(r"^_.+_$", s):      # whole-line italic meta note -> drop
            i += 1
            continue
        if lib_docx.is_table_row(s):
            block = []
            while i < n and lib_docx.is_table_row(lines[i].rstrip()):
                block.append(lines[i].rstrip())
                i += 1
            lib_docx.render_md_table(doc, block, add_inline)
            continue
        if s.startswith("## "):
            doc.add_heading(s[3:].strip(), level=2)
        elif s.startswith("# "):
            doc.add_heading(h1_title or s[2:].strip(), level=1)
        else:
            p = doc.add_paragraph()
            add_inline(p, s)
        i += 1

doc = Document()
# base style
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
for para_style in ("Normal",):
    doc.styles[para_style].paragraph_format.space_after = Pt(8)
    # Cardiovascular Diabetology instructs DOUBLE line spacing for the submitted manuscript
    # (round-4 Tier 1). 1.15 was a readability default, not the journal's rule.
    doc.styles[para_style].paragraph_format.line_spacing = 2.0

# ================= TITLE PAGE =================
# Read the title from the ONE place it is authored (FRONT_MATTER.md). It was previously a hardcoded
# literal here, which is a second copy of a fact -- exactly how this docx came to carry the retired
# title after FRONT_MATTER.md was updated to the locked one. build_submission_v1.py asserts the two
# agree; deriving it means they cannot disagree.
_front_src = open(os.path.join(FIG, "FRONT_MATTER.md"), encoding="utf-8").read()
_mt = re.search(r"##\s*Title\s*\n\*\*(.+?)\*\*", _front_src, flags=re.S)
assert _mt, "could not read the canonical title from FRONT_MATTER.md"
title = " ".join(_mt.group(1).split())
h = doc.add_paragraph(); h.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = h.add_run(title); r.bold = True; r.font.size = Pt(15)

doc.add_paragraph()
# Round-7 C002: the author list, affiliations and corresponding-author block are DERIVED from
# FRONT_MATTER.md - the same source the title page reads - because a hardcoded literal here is how
# this file came to ship "[Author list to be supplied]" on page 1 through six review rounds while a
# complete list sat in TITLE_PAGE.docx. gate_package asserts the two agree and that no bracketed
# placeholder survives.
_ma = re.search(r"##\s*Authors\s*\n(.+?)\n\s*\n", _front_src, flags=re.S)
assert _ma, "could not read the author list from FRONT_MATTER.md"
authors = doc.add_paragraph(); authors.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_inline(authors, " ".join(_ma.group(1).split()))

_maf = re.search(r"##\s*Affiliations\s*\n(.+?)\n##", _front_src, flags=re.S)
assert _maf, "could not read the affiliations from FRONT_MATTER.md"
for _line in [x.strip() for x in _maf.group(1).split("\n") if x.strip()]:
    _p = doc.add_paragraph(); _p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _r = _p.add_run(_line); _r.font.size = Pt(9)

_mc = re.search(r"Corresponding author: \*\*(.+?)\*\*\n(.*?)(?:\n\s*\n|\n##)", _front_src, flags=re.S)
assert _mc, "could not read the corresponding-author block from FRONT_MATTER.md"
_cblock = " ".join(_mc.group(2).split())
_mail = re.search(r"Email:\s*(\S+)", _cblock)
assert _mail, "the corresponding-author block carries no email address"
corr = doc.add_paragraph()
corr.add_run("Corresponding author: ").bold = True
corr.add_run(f"{_mc.group(1)}. {_cblock}")

# Running head and keywords are AUTHORED in FRONT_MATTER.md and read from there. They used to be
# literals here, i.e. second copies of two facts -- and the running head duly went stale when the
# title changed on 2026-07-25, still announcing the staging test the revision had demoted.
def _front_field(heading):
    m = re.search(r"##\s*" + re.escape(heading) + r"\s*\n(.+?)(?=\n##|\n---|\Z)",
                  _front_src, flags=re.S)
    assert m, f"FRONT_MATTER.md has no '## {heading}' section"
    return " ".join(m.group(1).replace("**", "").split())

meta = doc.add_paragraph()
meta.add_run("Running head: ").bold = True
meta.add_run(_front_field("Running head"))
# Round-7 C109: the journal places Keywords AFTER the Abstract and before Background, so they are
# emitted there (below), not in the title block where they used to share a paragraph with the
# running head.

# Round-4 Tier 1: the journal instructs NO manual page breaks. The section headings
# carry the structure instead.

# ================= ABSTRACT =================
doc.add_heading("Abstract", level=1)
abs_txt = read("ABSTRACT.md")
# pull the 4 structured paragraphs of the FULL abstract (between the FULL and CAP markers)
mfull = re.search(r"## FULL abstract.*?\n(.*?)\n## CAP variant", abs_txt, flags=re.S)
block = mfull.group(1) if mfull else abs_txt
for para in re.findall(r"\*\*(?:Background|Methods|Results|Conclusions)\.\*\*.*?(?=\n\n|\Z)",
                       block, flags=re.S):
    para = " ".join(para.split())   # unwrap
    p = doc.add_paragraph(); add_inline(p, para)

# ---- Keywords: BMC places them after the Abstract, before Background (round-7 C109) ----
_kw = doc.add_paragraph()
_kw.add_run("Keywords: ").bold = True
_kw.add_run(_front_field("Keywords"))

# ---- Research Insights: Cardiovascular Diabetology asks for this table "below the abstract" ----
ri = read("RESEARCH_INSIGHTS.md")
doc.add_heading("Research Insights", level=1)
_rows = [ln for ln in ri.split("\n") if ln.startswith("| ") and not re.match(r"^\|[-| ]+\|$", ln)]
_rows = [r for r in _rows if r.strip("| ").strip()]
# Drop the markdown header row: it is column titles, not a Research Insights question.
_rows = [r for r in _rows if not r.startswith("| Question |")]
_tbl = doc.add_table(rows=0, cols=2); _tbl.style = "Table Grid"
for _r in _rows:
    _cells = [c.strip() for c in _r.strip().strip("|").split("|")]
    if len(_cells) != 2:
        continue
    _row = _tbl.add_row().cells
    for _i, _c in enumerate(_cells):
        _row[_i].text = ""
        # Highlights are separated by <br> in the source so each is its own line in the cell:
        # the journal asks for discrete highlights, not a run-on paragraph.
        _parts = [x.strip() for x in _c.split("<br>") if x.strip()]
        add_inline(_row[_i].paragraphs[0], _parts[0] if _parts else "")
        for _extra in _parts[1:]:
            add_inline(_row[_i].add_paragraph(), _extra)
assert len(_tbl.rows) == 4, f"Research Insights table has {len(_tbl.rows)} rows, expected 4"

# ================= BODY =================
add_body(doc, read("INTRODUCTION.md"), h1_title="Background")
add_body(doc, read("METHODS.md"))
add_body(doc, read("RESULTS.md"))
add_body(doc, read("DISCUSSION.md"))
# Cardiovascular Diabetology requires Conclusions and a List of abbreviations as their own sections,
# in this order, between the Discussion and the Declarations.
add_body(doc, read("CONCLUSIONS.md"))
add_body(doc, read("ABBREVIATIONS.md"))

# ================= DECLARATIONS =================
front = read("FRONT_MATTER.md")
mdec = re.search(r"# Declarations(.*)$", front, flags=re.S)
if mdec:
    doc.add_heading("Declarations", level=1)
    # This loop, not add_body, is what renders the Additional-files table, and it had no table
    # branch: the block shipped as seven literal paragraphs, `|---|---|---|---|` included. Two
    # renderers in one file, only one of which was ever taught about tables.
    _dec = mdec.group(1).split("\n")
    _i, _n = 0, len(_dec)
    while _i < _n:
        s = _dec[_i].rstrip()
        if not s or s.startswith("#"):
            _i += 1
            continue
        if lib_docx.is_table_row(s):
            _blk = []
            while _i < _n and lib_docx.is_table_row(_dec[_i].rstrip()):
                _blk.append(_dec[_i].rstrip())
                _i += 1
            lib_docx.render_md_table(doc, _blk, add_inline)
            continue
        # Round-7 C051: accumulate consecutive wrapped lines into ONE paragraph. DECLARATIONS.md is
        # hard-wrapped at ~110 characters; emitting a paragraph per source line printed four pages of
        # broken fragments in the manuscript while the standalone DECLARATIONS.docx (render_md, which
        # unwraps) read correctly - a defect invisible to a twin-identity gate because the two
        # representations came from different renderers.
        _para = [s]
        _i += 1
        while _i < _n:
            _nxt = _dec[_i].rstrip()
            if not _nxt or _nxt.startswith("#") or lib_docx.is_table_row(_nxt):
                break
            _para.append(_nxt)
            _i += 1
        p = doc.add_paragraph(); add_inline(p, " ".join(_para))

# ================= REFERENCES (numbered Vancouver, by first appearance) =================
doc.add_heading("References", level=1)
# Vancouver: numbered, in order of first appearance. Rendered from the SAME ordering the in-text
# rewrite used (VAN.numbering()), so an in-text [n] and the nth entry cannot disagree.
for cite in VAN.reference_list(MASTER, REF_ORDER):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
    add_inline(p, cite)

# ================= TABLES =================
# Round-7 C014 / PI decision Q7: the graded evidence synthesis is a MAIN-TEXT table, generated by
# scripts/build_tables.py from the analysis outputs. It used to be a text panel inside a
# supplementary figure and a parallel list in the Results, which disagreed with each other.
_tables_src = read("TABLES.md")
doc.add_heading("Tables", level=1)
_tb = _tables_src.split("\n")
_ti, _tn = 0, len(_tb)
while _ti < _tn:
    _s = _tb[_ti].rstrip()
    if not _s or _s.startswith("# "):
        _ti += 1
        continue
    if lib_docx.is_table_row(_s):
        _blk = []
        while _ti < _tn and lib_docx.is_table_row(_tb[_ti].rstrip()):
            _blk.append(_tb[_ti].rstrip())
            _ti += 1
        lib_docx.render_md_table(doc, _blk, add_inline)
        continue
    _p = doc.add_paragraph(); add_inline(_p, _s)
    _ti += 1

# ================= FIGURE LEGENDS =================
add_body(doc, read("FIGURE_LEGENDS.md"), h1_title="Figure Legends")

# ================= line + page numbering (round-4 Tier 1) =================
# The journal instructs continuous line numbering and page numbers. python-docx exposes neither, so
# both are written directly into the section properties / footer as OOXML.
from docx.oxml.ns import qn as _qn
from docx.oxml import OxmlElement as _El

_sect = doc.sections[0]
_sectPr = _sect._sectPr
_ln = _El("w:lnNumType")
_ln.set(_qn("w:countBy"), "1")
_ln.set(_qn("w:start"), "1")
_ln.set(_qn("w:restart"), "continuous")
_sectPr.append(_ln)

_footer_p = _sect.footer.paragraphs[0]
_footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
_fld = _El("w:fldSimple")
_fld.set(_qn("w:instr"), "PAGE")
_run = _El("w:r"); _t = _El("w:t"); _t.text = "1"; _run.append(_t); _fld.append(_run)
_footer_p._p.append(_fld)

# ================= metadata (global rule 8) =================
cp = doc.core_properties
cp.author = "Bertrand Chin-Ming Tan"
cp.last_modified_by = "Bertrand Chin-Ming Tan"
cp.title = title
cp.comments = ""
# Otherwise the file advertises python-docx's 2013-12-23 template stamp as its creation date.
import datetime as _dt
cp.created = cp.modified = _dt.datetime.now()

path = os.path.join(OUT, "CKM_Paper4_manuscript.docx")
doc.save(path)
# quick counts
print(f"wrote {path}")
print(f"references rendered: {len(REF_ORDER)} (numbered Vancouver, by first appearance)")
print(f"paragraphs: {len(doc.paragraphs)}")
